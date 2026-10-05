import csv, io, re
from docx import Document as DocxDocument

def text_txt(text): return text.encode("utf-8")
def text_docx(text):
    doc=DocxDocument()
    for line in text.splitlines(): doc.add_paragraph(line)
    out=io.BytesIO(); doc.save(out); return out.getvalue()
def text_csv(text):
    out=io.StringIO(); w=csv.writer(out); w.writerow(["line","text"])
    for i,line in enumerate(text.splitlines(),1): w.writerow([i,line])
    return out.getvalue().encode("utf-8")


def table_xlsx(rows, headers):
    from openpyxl import Workbook
    out=io.BytesIO(); wb=Workbook(); ws=wb.active
    ws.append(headers)
    for row in rows: ws.append(list(row))
    wb.save(out); return out.getvalue()
def table_pdf(rows, headers, title="Report"):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    out=io.BytesIO(); doc=SimpleDocTemplate(out,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
    styles=getSampleStyleSheet(); data=[headers]+[[str(x) if x is not None else "" for x in row] for row in rows]
    t=Table(data,repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#152238")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
                           ("GRID",(0,0),(-1,-1),0.25,colors.grey),("FONTSIZE",(0,0),(-1,-1),7),("VALIGN",(0,0),(-1,-1),"TOP")]))
    doc.build([Paragraph(title,styles["Heading2"]),t]); return out.getvalue()
